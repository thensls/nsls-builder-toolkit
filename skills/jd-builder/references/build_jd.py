"""Render an NSLS job description by filling the official template.

USAGE
  ~/.local/bin/nsls-python build_jd.py <spec.json> [--out-dir DIR]   (Windows: nsls-python ...)

The spec is JSON:
  {
    "title": "Account Manager",                 # required
    "location": "Remote",                       # optional, default "Remote"
    "purpose": "keeping every chapter ...",     # finishes "The <title> will be responsible for ..."
    "team": "Client Services team, which ...",  # finishes "This role forms part of the ..."
    "responsibilities": ["Lead ...", ...],      # <= 10, same verb tense
    "qualifications": ["...", ...],             # <= 10, objective hard skills
    "nice_to_haves": ["...", ...],
    "who_you_are": ["...", ...],                # per-role traits
    "how_we_work": [...],                       # optional; default = references/how-we-work.md
    "status": "DRAFT"                           # DRAFT (default) or FINAL
  }

It fills ONLY the placeholders. The logo, the opener, the guiding question and
the EEO line are kept byte-for-byte. Prints the output path; the SKILL.md tells
you how to upload it.
"""
import copy
import json
import os
import re
import sys

sys.path.insert(0, os.path.expanduser('~/.local/lib/nsls-pydeps'))  # toolkit's durable deps home
from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "jd-template-2025.docx")
HOW_WE_WORK = os.path.join(HERE, "how-we-work.md")

MAX_BULLETS = {"responsibilities": 10, "qualifications": 10}


def fail(msg):
    sys.exit(f"build_jd: {msg}")


def default_how_we_work():
    with open(HOW_WE_WORK, encoding="utf-8") as f:
        return [ln[2:].strip() for ln in f if ln.startswith("- ")]


def is_list(p):
    return p._p.find(".//" + qn("w:numPr")) is not None


def set_list(paragraphs, items):
    """Replace a run of template bullet paragraphs with one bullet per item."""
    if not paragraphs:
        fail("template changed: expected bullet paragraphs not found")
    proto = paragraphs[0]._p
    anchor = proto
    for item in items:
        new = copy.deepcopy(proto)
        runs = new.findall(qn("w:r"))
        for r in runs[1:]:
            new.remove(r)
        if runs:
            for t in runs[0].findall(qn("w:t")):
                runs[0].remove(t)
            t = runs[0].makeelement(qn("w:t"), {})
            t.text = item
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            runs[0].append(t)
        else:
            fail("template bullet has no run to copy formatting from")
        anchor.addnext(new)
        anchor = new
    for p in paragraphs:
        p._p.getparent().remove(p._p)


def bullets_after(doc, heading):
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip().rstrip(":").strip().lower() == heading.lower():
            out = []
            for q in paras[i + 1:]:
                if not is_list(q):
                    break
                out.append(q)
            return out
    fail(f"template changed: heading '{heading}' not found")


def replace_run(p, old, new):
    hit = False
    for r in p.runs:
        if r.text == old or r.text.startswith(old):
            r.text = r.text.replace(old, new, 1) if r.text != old else new
            hit = True
    if not hit:
        fail(f"template changed: placeholder '{old[:40]}' not found")


def main():
    if len(sys.argv) < 2:
        fail(__doc__)
    spec_path = sys.argv[1]
    out_dir = os.path.dirname(os.path.abspath(spec_path))
    if "--out-dir" in sys.argv:
        out_dir = sys.argv[sys.argv.index("--out-dir") + 1]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)

    for k in ("title", "purpose", "team", "responsibilities", "qualifications", "who_you_are"):
        if not spec.get(k):
            fail(f"spec missing '{k}'")
    for k, n in MAX_BULLETS.items():
        if len(spec[k]) > n:
            fail(f"'{k}' has {len(spec[k])} bullets; the template caps it at {n}")
    status = spec.get("status", "DRAFT").upper()
    if status == "FINAL":
        # Check the text the Doc will show, not json.dumps(spec): serialized
        # lists are themselves wrapped in [ ], so that check rejected every FINAL.
        texts = [v for v in spec.values() if isinstance(v, str)]
        texts += [s for v in spec.values() if isinstance(v, list) for s in v if isinstance(s, str)]
        open_brackets = [t for t in texts if re.search(r"\[[^\]]+\]", t)]
        if open_brackets:
            fail(f"status FINAL but [brackets] are still open: {open_brackets[:3]}")

    title = spec["title"].strip()
    doc = Document(TEMPLATE)
    paras = doc.paragraphs

    title_p = next(p for p in paras if p.text.startswith("(Title)"))
    replace_run(title_p, "(Title) ", f"{title} ")
    if spec.get("location"):
        replace_run(title_p, "The National Society of Leadership and Success - Remote",
                    f"The National Society of Leadership and Success - {spec['location']}")

    ov = next(p for p in paras if p.text.startswith("Overview:"))
    for r in ov.runs:
        if r.text == "(title)":
            r.text = title
        elif r.text.startswith("(high level overview"):
            r.text = spec["purpose"].strip().rstrip(".") + ". "

    team_p = next(p for p in paras if p.text.startswith("This role forms part of the"))
    for r in team_p.runs:
        if r.text.startswith("(insert team"):
            r.text = spec["team"].strip() + " "

    cell = doc.tables[0].cell(0, 0)
    set_list([p for p in cell.paragraphs if is_list(p)], spec["responsibilities"])

    set_list(bullets_after(doc, "Qualifications"), spec["qualifications"])
    set_list(bullets_after(doc, "Nice To Haves"), spec.get("nice_to_haves") or ["[none]"])
    set_list(bullets_after(doc, "Who You Are"), spec["who_you_are"])
    set_list(bullets_after(doc, "How We Work"), spec.get("how_we_work") or default_how_we_work())

    leftover = [p.text for p in doc.paragraphs if re.search(r"\((title|insert|include|high level|this )", p.text, re.I)]
    leftover += [p.text for p in cell.paragraphs if p.text.startswith("(")]
    if leftover:
        fail(f"template guidance text left in the Doc: {leftover[:2]}")

    safe = re.sub(r'[\\/:*?"<>|]', "-", title)
    out = os.path.join(out_dir, f"JD - {safe} ({status}).docx")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
